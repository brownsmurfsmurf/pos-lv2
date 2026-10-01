# -*- coding: utf-8 -*-
"""Markdown → docx の簡易変換（見出し・段落・箇条書き・表・コード・引用・画像）。

使い方: python tools/md2docx.py 入力.md 出力.docx
画像は段落として単独で書かれた ![alt](path) のみ埋め込む（表のセル内は不可）。
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

FONT = "Meiryo"
MONO = "Consolas"


def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT if name == MONO else name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text, size=10.5, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), name=MONO, size=size - 0.5, bold=bold, color=(0x8B, 0x1A, 0x1A))
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            set_font(par.add_run(m.group(1)), size=size, bold=bold, color=(0x2F, 0x55, 0x97))
        else:
            set_font(par.add_run(part), size=size, bold=bold)


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[2:]
    ncol = len(header)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ""
        add_inline(cell.paragraphs[0], text, size=9.5, bold=True)
        shade(cell, "DEEBF7")
    for r in body:
        r = (r + [""] * ncol)[:ncol]
        cells = table.add_row().cells
        for i, text in enumerate(r):
            cells[i].paragraphs[0].text = ""
            add_inline(cells[i].paragraphs[0], text.replace("<br>", "\n"), size=9.5)
    doc.add_paragraph()


def add_code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Mm(4)
        set_font(p.add_run(ln if ln else " "), name=MONO, size=9)
    doc.add_paragraph()


def convert(src, dst):
    base = os.path.dirname(os.path.abspath(src))
    lines = io.open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = sec.bottom_margin = Mm(20)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10.5)

    i, first_h2 = 0, True
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(doc, code)
        elif s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^#{1,4} ", s):
            level = len(s) - len(s.lstrip("#"))
            text = s[level:].strip()
            if level == 2 and not first_h2:
                doc.add_page_break()
            if level == 2:
                first_h2 = False
            p = doc.add_heading(level=min(level, 4))
            size = {1: 18, 2: 14, 3: 12, 4: 11}[min(level, 4)]
            add_inline(p, text, size=size, bold=True)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x29, 0x33)
        elif re.match(r"^!\[[^\]]*\]\([^)]+\)$", s):
            path = os.path.join(base, re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", s).group(1))
            if os.path.isfile(path):
                doc.add_picture(path, width=Mm(165))
        elif s in ("---", "***"):
            pass
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(6)
            add_inline(p, s.lstrip("> ").strip(), size=10.5)
            for r in p.runs:
                if r.font.color.rgb is None:
                    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        elif re.match(r"^\s*[-*] ", line):
            indent = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Mm(6 + 6 * indent)
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\s*\d+\. ", "", line))
        elif s == "":
            pass
        else:
            para = [s]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(
                r"^(\s*[-*] |\s*\d+\. |#{1,4} |\||```|>|!\[|---)", lines[i + 1].strip()
            ):
                i += 1
                para.append(lines[i].strip())
            add_inline(doc.add_paragraph(), "".join(para))
        i += 1
    doc.save(dst)
    print("written", dst, os.path.getsize(dst), "bytes")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: python tools/md2docx.py in.md out.docx")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
